"""Downloadable, editable user-reviewed case document drafts."""

from io import BytesIO

from docx import Document


def _lines(result, title):
    fields = result.document_fields
    get = lambda key: str(fields.get(key) or "【请填写】")
    return [
        title,
        "（以下为草稿，请逐项核对后再发送）",
        "致：" + get("school_name"),
        "本人" + get("student_name") + "于" + get("contract_date") + "报名参加驾驶培训，已支付" + get("amount") + "。",
        "现就培训服务争议说明如下：",
        result.fact_summary or "【请补充事实经过】",
        "诉求：" + str(fields.get("request") or "【请填写具体诉求及金额计算依据】"),
        "协商建议：" + (result.negotiation or "【请填写】"),
        "法律依据（请核对适用性）：",
        *[f"- {law.get('title', '【请核对】')} {law.get('article', '')}：{law.get('source_url', '')}" for law in result.legal_basis],
        "建议渠道：",
        *[f"- {route.get('channel', '')}：{route.get('when_to_use', '')}" for route in result.routes],
        "待核实事项：" + ("；".join(result.uncertainties) if result.uncertainties else "无"),
        "证据清单：培训合同、付款凭证、课时记录、沟通记录及事件时间线（按实际情况附上）。",
        "请在此填写日期、联系方式并签名：________________",
    ]


def export_document(result, format, title="协商函"):
    lines = _lines(result, title)
    if format == "txt":
        return "\n".join(lines).encode("utf-8")
    if format != "docx":
        raise ValueError("format must be docx or txt")
    doc = Document()
    doc.add_heading(title, level=1)
    for line in lines[1:]:
        doc.add_paragraph(line)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
