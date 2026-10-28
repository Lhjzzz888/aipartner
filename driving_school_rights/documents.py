"""In-memory evidence extraction and redaction helpers."""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader
import pypdfium2 as pdfium
import pytesseract


class DocumentError(ValueError):
    """Raised when an evidence file cannot be safely processed."""


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str
    file_type: str
    text: str
    warnings: tuple[str, ...] = ()
    page_count: int | None = None
    ocr_confidence: float | None = None


def extract_document(filename, content, config):
    suffix = Path(filename).suffix.lower()
    supported = {".pdf", ".docx", ".txt", ".png", ".jpg", ".jpeg"}
    if suffix not in supported:
        raise DocumentError("暂不支持此文件类型，请上传 PDF、DOCX、TXT、PNG 或 JPG。")
    if not isinstance(content, bytes) or not content:
        raise DocumentError("文件为空或格式无效。")
    if len(content) > config.max_upload_mb * 1024 * 1024:
        raise DocumentError(f"文件大小超过 {config.max_upload_mb} MB 上限。")

    warnings = []
    page_count = None
    confidence = None
    try:
        if suffix == ".pdf":
            if not content.startswith(b"%PDF"):
                raise DocumentError("PDF 文件格式无效或已损坏。")
            reader = PdfReader(BytesIO(content), strict=False)
            page_count = len(reader.pages)
            if page_count > config.max_pdf_pages:
                raise DocumentError(f"PDF 页数超过 {config.max_pdf_pages} 页上限，请拆分后上传。")
            chunks = [""] * page_count
            scanned_pages = []
            for number, page in enumerate(reader.pages):
                page_text = (page.extract_text() or "").strip()
                if page_text:
                    chunks[number] = page_text
                else:
                    scanned_pages.append(number)
            if scanned_pages:
                rendered = pdfium.PdfDocument(content)
                try:
                    for number in scanned_pages:
                        bitmap = rendered[number].render(scale=2)
                        image = bitmap.to_pil().copy()
                        ocr_text, page_confidence = _ocr(image)
                        if ocr_text:
                            chunks[number] = ocr_text
                        if page_confidence is not None:
                            confidence = page_confidence if confidence is None else min(confidence, page_confidence)
                finally:
                    rendered.close()
                warnings.append("已对无文本层的 PDF 页面执行 OCR，请核对识别结果。")
        elif suffix == ".docx":
            if not content.startswith(b"PK"):
                raise DocumentError("DOCX 文件格式无效或已损坏。")
            doc = Document(BytesIO(content))
            chunks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        chunks.append(" | ".join(cells))
        elif suffix == ".txt":
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                try:
                    text = content.decode("gb18030")
                except UnicodeDecodeError as exc:
                    raise DocumentError("TXT 编码无法识别，请保存为 UTF-8 或 GB18030 后重试。") from exc
            chunks = [text.strip()]
        else:
            try:
                with Image.open(BytesIO(content)) as source:
                    image_format = source.format
                    source.verify()
                expected_format = "PNG" if suffix == ".png" else "JPEG"
                if image_format != expected_format:
                    raise DocumentError("图片实际格式与文件扩展名不符，请重新导出为 PNG 或 JPG。")
                with Image.open(BytesIO(content)) as source:
                    image = source.convert("RGB")
            except DocumentError:
                raise
            except (UnidentifiedImageError, OSError) as exc:
                raise DocumentError("图片格式无效或已损坏。") from exc
            text, confidence = _ocr(image)
            chunks = [text]
            warnings.append("已对图片执行 OCR，请核对识别结果。")
    except DocumentError:
        raise
    except Exception as exc:
        if suffix == ".pdf":
            raise DocumentError("PDF 无法读取，请检查文件是否损坏或加密。") from exc
        if suffix == ".docx":
            raise DocumentError("DOCX 无法读取，请检查文件是否损坏或加密。") from exc
        raise DocumentError("文件处理失败，请检查文件内容后重试。") from exc

    text = "\n".join(part for part in chunks if part).strip()
    if confidence is not None and confidence < 75:
        warnings.append("OCR 识别置信度偏低，请对照原件人工校正。")
    if not text:
        warnings.append("没有提取到可用文字；请尝试更清晰的扫描件或手动补充事实。")
    if len(text) > config.max_text_chars:
        text = text[: config.max_text_chars]
        warnings.append(f"提取文字已截断至 {config.max_text_chars} 字符长度上限。")
    return ExtractedDocument(
        filename=Path(filename.replace("\\", "/")).name,
        file_type=suffix[1:],
        text=text,
        warnings=tuple(warnings),
        page_count=page_count,
        ocr_confidence=confidence,
    )


def _ocr(image):
    try:
        data = pytesseract.image_to_data(
            image, lang="chi_sim+eng", output_type=pytesseract.Output.DICT
        )
    except pytesseract.TesseractNotFoundError as exc:
        raise DocumentError("服务器未安装 Tesseract OCR，暂时无法识别扫描件。") from exc
    except pytesseract.TesseractError as exc:
        detail = str(exc).lower()
        if "chi_sim" in detail or "failed loading language" in detail:
            raise DocumentError("服务器缺少简体中文 OCR 语言包，请检查部署依赖。") from exc
        raise DocumentError("OCR 处理失败，请更换清晰图片后重试。") from exc
    words = [word.strip() for word in data.get("text", []) if word.strip()]
    scores = []
    for value in data.get("conf", []):
        try:
            score = float(value)
        except (TypeError, ValueError):
            continue
        if score >= 0:
            scores.append(score)
    confidence = round(sum(scores) / len(scores), 1) if scores else None
    return " ".join(words), confidence


def sanitize_text(text, redactions):
    for value in sorted((str(item) for item in redactions if str(item)), key=len, reverse=True):
        text = text.replace(value, "[已遮盖]")
    return text
