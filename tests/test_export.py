from docx import Document
from io import BytesIO

from driving_school_rights.deepseek import AssistantResult
from driving_school_rights.export import export_document


def _result():
    return AssistantResult("驾校尚未安排课程", ["已付金额？"], [{"title":"民法典", "article":"第五百七十七条", "source_url":"https://example.gov/law"}], "先书面协商", [{"channel":"12328", "when_to_use":"驾培管理事项"}], {"student_name":"张三", "school_name":"某驾校", "amount":"2000元"}, ["退费金额待核对"])


def test_export_docx_has_confirmed_fields_and_sources():
    data = export_document(_result(), "docx", "协商函")
    doc = Document(BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "张三" in text and "某驾校" in text and "https://example.gov/law" in text


def test_export_keeps_unknown_fields_as_placeholders():
    result = _result()
    result.document_fields["student_name"] = ""
    text = export_document(result, "txt", "协商函").decode("utf-8")
    assert "【请填写】" in text


def test_export_txt_is_utf8():
    text = export_document(_result(), "txt", "协商函").decode("utf-8")
    assert "协商函" in text and "退费金额待核对" in text
